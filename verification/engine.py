"""Verify that a page represents published SIH 2026 screening records."""
from __future__ import annotations
import re
from dataclasses import dataclass
from bs4 import BeautifulSoup

@dataclass(frozen=True)
class VerificationResult:
    verified: bool
    edition: str | None
    reason: str
    headers: tuple[str,...] = ()
    record_count: int = 0

class VerificationEngine:
    EDITION=re.compile(r"SIH\s*2026|Smart\s+India\s+Hackathon\s*2026",re.I)
    PLACEHOLDER=re.compile(r"(coming soon|not yet published|no results? available|under construction)",re.I)
    RESULT=re.compile(r"(screening results?|shortlisted teams|selected teams)",re.I)
    HEADER_GROUPS=(
        {"team name","team"}, {"status","screening status"},
        {"aishe code","institution code","institute code"},
        {"name of the institute","institute name","institution name","college name"},
    )

    def verify_html(self, html: str, url: str = "") -> VerificationResult:
        soup=BeautifulSoup(html,"html.parser")
        text=soup.get_text(" ",strip=True)
        if not self.EDITION.search(text): return VerificationResult(False,None,"SIH 2026 edition marker missing")
        if self.PLACEHOLDER.search(text): return VerificationResult(False,"SIH 2026","placeholder/unpublished marker found")
        if not self.RESULT.search(text): return VerificationResult(False,"SIH 2026","result/screening heading missing")
        tables=soup.find_all("table")
        for table in tables:
            rows=table.find_all("tr")
            if not rows: continue
            headers=tuple(c.get_text(" ",strip=True).lower() for c in rows[0].find_all(["th","td"]))
            joined=set(headers)
            groups=sum(bool(group & joined) for group in self.HEADER_GROUPS)
            data_rows=[r for r in rows[1:] if len(r.find_all(["td","th"])) >= max(3,len(headers)-2)]
            if groups >= 3 and data_rows:
                return VerificationResult(True,"SIH 2026","edition, result context, expected fields, and records found",headers,len(data_rows))
        return VerificationResult(False,"SIH 2026","no populated result table with expected fields")

    def verify_document(self, data: bytes, content_type: str = "", filename: str = "", url: str = "") -> VerificationResult:
        """Verify a directly published HTML, CSV, PDF, or Excel result document."""
        name=filename.casefold(); mime=content_type.casefold()
        if data[:5] != b"%PDF-" and not name.endswith((".csv", ".pdf", ".xls", ".xlsx")) and not any(
            hint in mime for hint in ("text/csv", "application/pdf", "spreadsheetml", "ms-excel")
        ):
            return self.verify_html(data.decode("utf-8", errors="replace"), url)
        try:
            if data[:5] == b"%PDF-" or name.endswith(".pdf") or "application/pdf" in mime:
                import fitz
                doc=fitz.open(stream=data,filetype="pdf")
                try: content=" ".join(page.get_text("text") for page in doc)
                finally: doc.close()
            elif name.endswith(".xlsx") or "spreadsheetml" in mime:
                from openpyxl import load_workbook
                book=load_workbook(__import__("io").BytesIO(data), read_only=True, data_only=True)
                try: content=" ".join(str(cell) for sheet in book.worksheets for row in sheet.iter_rows(values_only=True) for cell in row if cell is not None)
                finally: book.close()
            elif name.endswith(".xls") or "ms-excel" in mime:
                import xlrd
                book=xlrd.open_workbook(file_contents=data, on_demand=True)
                content=" ".join(str(value) for sheet in book.sheets() for row in range(sheet.nrows) for value in sheet.row_values(row))
                book.release_resources()
            else:
                content=data.decode("utf-8-sig", errors="replace")
        except Exception as exc:
            return VerificationResult(False, None, f"could not read result document: {exc}")
        text=f"{url} {filename} {content}"
        if not self.EDITION.search(text):
            return VerificationResult(False,None,"SIH 2026 edition marker missing")
        if self.PLACEHOLDER.search(text):
            return VerificationResult(False,"SIH 2026","placeholder/unpublished marker found")
        # Reuse the extractor's alias table as a schema check so editions can vary
        # in header wording without accepting arbitrary documents.
        from extraction.engine import ExtractionEngine, ExtractionError
        try:
            extracted=ExtractionEngine().extract(data,content_type,filename)
        except ExtractionError as exc:
            return VerificationResult(False,"SIH 2026",f"no recognizable result records: {exc}")
        if not self.RESULT.search(text) and not re.search(r"screening.?status|selection.?status|shortlist|waitlist", text, re.I):
            return VerificationResult(False,"SIH 2026","result/screening context missing")
        return VerificationResult(True,"SIH 2026","SIH 2026 result document and populated record schema found",
                                  tuple(extracted.headers),len(extracted.records))
