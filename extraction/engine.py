"""Extract tabular SIH records from HTML, CSV, Excel, and text-based PDF files."""
from __future__ import annotations
import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from typing import Any
from bs4 import BeautifulSoup

ALIASES={
 "problem_statement_id":("ps id","psid","problem statement id","problem statement","problem id"),
 "team_name":("team name","name of team","team"),
 "team_id":("team id","sih team id"),
 "idea_id":("idea id","idea number"),
 "institution_code":("aishe code","aishe id","institution code","institute code","college code"),
 "institution_name":("name of the institute","institute name","institution name","name of institute","college name","institute/college name"),
 "city":("institute city","institute district","city","district"),
 "state":("institute state","state"),
 "status":("status","screening status","result","selection status","team status"),
 "team_leader":("team leader name","leader name","team lead","team leader"),
}

def normalize_text(value: Any) -> str:
    value=unicodedata.normalize("NFKD",str(value or ""))
    value="".join(ch for ch in value if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+"," ",value.casefold()).strip()

def _column_map(headers: list[str]) -> dict[str,int]:
    normalized=[normalize_text(h) for h in headers]
    mapping={}
    for field,aliases in ALIASES.items():
        accepted={normalize_text(alias) for alias in aliases}
        for i,header in enumerate(normalized):
            if header in accepted:
                mapping[field]=i; break
    return mapping

@dataclass
class ExtractedBatch:
    records: list[dict[str,Any]]
    headers: list[str]
    source_format: str

class ExtractionError(ValueError): pass

class ExtractionEngine:
    REQUIRED={"team_name","institution_name","status"}

    def from_html(self, html: str) -> ExtractedBatch:
        soup=BeautifulSoup(html,"html.parser")
        all_records=[]; first_headers=[]
        for table in soup.find_all("table"):
            rows=table.find_all("tr")
            if not rows: continue
            first_cells=rows[0].find_all(["th","td"])
            headers=[c.get_text(" ",strip=True) for c in first_cells]
            mapping=_column_map(headers)
            if not self.REQUIRED.issubset(mapping): continue
            if not first_headers: first_headers=headers
            records=[]
            for row in rows[1:]:
                cells=row.find_all(["td","th"])
                if len(cells)<len(headers): continue
                values=[c.get_text(" ",strip=True) for c in cells[:len(headers)]]
                if not any(values): continue
                records.append(self._record(headers,values,mapping))
            all_records.extend(records)
        if all_records:
            # Multiple result tables are common when publications split divisions
            # or statuses. Keep every qualifying table; pipeline fingerprints dedupe.
            return ExtractedBatch(all_records,first_headers,"html")
        raise ExtractionError("No populated HTML table had recognizable team, institution, and status columns")

    def from_csv(self, data: bytes, encoding: str="utf-8-sig") -> ExtractedBatch:
        text=data.decode(encoding,errors="replace")
        try: dialect=csv.Sniffer().sniff(text[:4096],delimiters=",;\t|")
        except csv.Error: dialect=csv.excel
        rows=list(csv.reader(io.StringIO(text),dialect))
        return self._from_rows(rows,"csv")

    def from_excel(self, data: bytes, filename: str) -> ExtractedBatch:
        suffix=filename.lower().rsplit(".",1)[-1]
        if suffix=="xlsx":
            from openpyxl import load_workbook
            book=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
            try:
                for sheet in book.worksheets:
                    result=self._try_rows(list(sheet.iter_rows(values_only=True)),"xlsx")
                    if result: return result
            finally: book.close()
        elif suffix=="xls":
            import xlrd
            book=xlrd.open_workbook(file_contents=data,on_demand=True)
            for sheet in book.sheets():
                result=self._try_rows([sheet.row_values(i) for i in range(sheet.nrows)],"xls")
                if result: return result
        else: raise ExtractionError(f"Unsupported Excel extension: {suffix}")
        raise ExtractionError("No worksheet contained recognizable result columns and rows")

    def from_pdf(self, data: bytes) -> ExtractedBatch:
        try: import fitz
        except ImportError as exc: raise ExtractionError("PDF support requires PyMuPDF") from exc
        doc=fitz.open(stream=data,filetype="pdf")
        try:
            rows=[]
            for page in doc:
                if hasattr(page,"find_tables"):
                    for table in page.find_tables().tables:
                        rows.extend(table.extract())
            result=self._try_rows(rows,"pdf") if rows else None
            if result: return result
            # Fallback for selectable-text PDFs with consistently spaced columns.
            lines=[]
            for page in doc:
                lines.extend(re.split(r"\s{2,}",line.strip()) for line in page.get_text("text").splitlines() if line.strip())
            result=self._try_rows(lines,"pdf")
            if result: return result
            raise ExtractionError("PDF contains no recognizable table; scanned PDFs require OCR/manual review")
        finally: doc.close()

    def extract(self, data: bytes, content_type: str="", filename: str="") -> ExtractedBatch:
        name=filename.lower(); mime=content_type.lower()
        if data[:5]==b"%PDF-" or name.endswith(".pdf") or "application/pdf" in mime: return self.from_pdf(data)
        if name.endswith(".xlsx") or "spreadsheetml" in mime: return self.from_excel(data,filename or "results.xlsx")
        if name.endswith(".xls") or "ms-excel" in mime: return self.from_excel(data,filename or "results.xls")
        if name.endswith(".csv") or "text/csv" in mime: return self.from_csv(data)
        try: return self.from_html(data.decode("utf-8",errors="replace"))
        except ExtractionError:
            try: return self.from_csv(data)
            except Exception as exc: raise ExtractionError("Unknown result file format or unsupported layout") from exc

    def _from_rows(self, rows: list[list[Any]], source_format: str) -> ExtractedBatch:
        result=self._try_rows(rows,source_format)
        if not result: raise ExtractionError("No rows contained recognizable team, institution, and status columns")
        return result

    def _try_rows(self, rows: list[Any], source_format: str) -> ExtractedBatch | None:
        for header_idx,row in enumerate(rows[:10]):
            headers=["" if x is None else str(x).strip() for x in row]
            mapping=_column_map(headers)
            if not self.REQUIRED.issubset(mapping): continue
            records=[]
            for raw in rows[header_idx+1:]:
                values=["" if x is None else str(x).strip() for x in raw]
                if len(values)<len(headers) or not any(values): continue
                records.append(self._record(headers,values,mapping))
            if records: return ExtractedBatch(records,headers,source_format)
        return None

    @staticmethod
    def _record(headers: list[str], values: list[str], mapping: dict[str,int]) -> dict[str,Any]:
        source={headers[i]:values[i] for i in range(min(len(headers),len(values)))}
        normalized={field:(values[idx].strip() if idx<len(values) else "") for field,idx in mapping.items()}
        return {"source":source,"normalized":normalized}
